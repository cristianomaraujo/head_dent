"""Read the reviewed DART workbook; collect conversations without assigning clinical scores.

prepare performs no network calls. run is explicitly bounded, checkpoints each turn,
and never silently retries a possibly billed request. export produces Excel-readable CSV.
"""
import argparse
from contextlib import contextmanager
import csv
from datetime import date, datetime, timezone
import hashlib
from importlib.metadata import version
import json
import os
import platform
from pathlib import Path
import sqlite3
import time

import httpx


SCHEMA = 1
FLOW = "site-triage-then-chat-v2"
GUIDELINES = {1: "ESE", 2: "AAE"}
SOURCE_ROOT = Path(__file__).resolve().parents[1]
METADATA_FIELDS = ("execution_id", "case_num", "case_id", "template_id", "trauma_code", "versao_code",
                   "diretriz_code", "idioma_code", "repeticao", "ordem_execucao", "analise_principal_code")


def initial_message(language="pt"):
    # Same message sent by app/static/index.html when the clinician starts a case.
    return ("Response language selected by the clinician: " + language + ". Respond only in that language. "
            "Receba esta ficha de triagem e faça a primeira resposta ao dentista. Se faltarem dados clínicos necessários, "
            "diga quais são e pergunte ao profissional, sem presumir os achados. Se os dados forem suficientes, "
            "apresente a orientação com base no DART.")


def encode(value):
    def dates(v):
        if isinstance(v, (date, datetime)):
            return v.isoformat()
        raise TypeError(f"Unsupported data: {type(v).__name__}")
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=dates)


def digest(value):
    return hashlib.sha256(value).hexdigest()


def now():
    return datetime.now(timezone.utc).isoformat()


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(str(value)).isoformat()


def table(book, name):
    rows = book[name].iter_rows(values_only=True)
    headers = list(next(rows))
    if any(not isinstance(h, str) or not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"Invalid or duplicate headers in {name}")
    return headers, [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]


def prepare(source, destination, reviewed_by, review_note=""):
    from openpyxl import load_workbook  # read-only extraction; never rewrites the source
    if not reviewed_by.strip():
        raise ValueError("Identify the person who confirmed the case/reference review")
    if destination.exists() and any(destination.iterdir()):
        raise ValueError("Use a new empty output directory; existing plans are immutable")
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        _, case_rows = table(book, "Casos")
        _, gold_rows = table(book, "Gabaritos")
        result_headers, result_rows = table(book, "Resultados")
    finally:
        book.close()
    if not set(METADATA_FIELDS).issubset(result_headers):
        raise ValueError("Missing planned execution columns in Resultados")
    cases = {c["case_id"]: c for c in case_rows}
    gold = {(g["case_id"], g["diretriz_code"]): g for g in gold_rows}
    if len(cases) != len(case_rows) or len(gold) != len(gold_rows):
        raise ValueError("Duplicate case or reference keys")
    catalog = json.loads((SOURCE_ROOT / "knowledge/catalog.json").read_text(encoding="utf-8"))
    options = {c["id"]: {q["key"]: {o["value"] for o in q["options"]} for q in c["questions"]} for c in catalog}
    conditions = {c["id"]: c for c in catalog}
    plan, seen_ids, seen_order, seen_configuration = [], set(), set(), set()
    required = ["caso_completo", "caso_dados_faltantes", "resposta_complementar", "dados_a_solicitar"]
    for c in case_rows:
        if any(not isinstance(c.get(k), str) or not c[k].strip() for k in required):
            raise ValueError(f"Missing scenario data: {c['case_id']}")
        if not all((c["case_id"], g) in gold for g in GUIDELINES):
            raise ValueError(f"Missing ESE/AAE references: {c['case_id']}")
    for r in result_rows:
        c = cases.get(r["case_id"])
        v, g, rep = r["versao_code"], r["diretriz_code"], r["repeticao"]
        key = (r["case_id"], v, g, rep)
        if c is None or v not in (1, 2) or g not in GUIDELINES or r["idioma_code"] != 1:
            raise ValueError(f"Invalid configuration: {r['execution_id']}")
        if not isinstance(rep, int) or rep < 1 or not isinstance(r["ordem_execucao"], int):
            raise ValueError("Repetition and execution order must be integers")
        if r["execution_id"] in seen_ids or r["ordem_execucao"] in seen_order or key in seen_configuration:
            raise ValueError("Duplicate execution ID, order, or configuration")
        if any(r[k] != c[k] for k in ("case_num", "case_id", "template_id", "trauma_code")):
            raise ValueError("Execution metadata does not match its parent case")
        # This collector never carries pre-existing clinical scores into a new study.
        derived = {"adequacao_clinica", "adequacao_operacional", "fracao_completude", "principal_adjudicado"}
        observations = set(result_headers) - set(METADATA_FIELDS) - derived - {"status_execucao"}
        if r.get("status_execucao") not in (None, 0) or any(r.get(k) not in (None, "") for k in observations):
            raise ValueError("Source contains observations; preserve them and use a clean collection plan")
        seen_ids.add(r["execution_id"]); seen_order.add(r["ordem_execucao"]); seen_configuration.add(key)
        injury = f"T{c['trauma_code']:02d}"
        answers = json.loads(c["triagem_completa_json" if v == 1 else "triagem_incompleta_json"])
        if injury not in options or not isinstance(answers, dict) or any(
                k not in options[injury] or value not in options[injury][k] for k, value in answers.items()):
            raise ValueError(f"Triage codes do not match the catalog: {c['case_id']}")
        context = {"injury_id": injury, "guideline": GUIDELINES[g], "answers": answers,
                   "case_text": "",
                   "reference_date": iso(c["data_evento"]),
                   "evaluation_date": iso(c["data_avaliacao"]), "language": "pt"}
        complete_answers = json.loads(c["triagem_completa_json"])
        triage = []
        for q in conditions[injury]["questions"]:
            labels = {o["value"]: o["label"] for o in q["options"]}
            selected = answers.get(q["key"])
            triage.append({"key": q["key"], "field": q["label"], "selected_code": selected,
                           "selected_label": labels.get(selected, "Não informado"),
                           "withheld": selected is None and q["key"] in complete_answers,
                           "reference_code": complete_answers.get(q["key"]),
                           "reference_label": labels.get(complete_answers.get(q["key"]), "")})
        turns = [
            {"message": initial_message(), "expected_status": None},
            {"message": "Acrescentando os dados clínicos do caso: " + c["caso_completo"],
             "expected_status": "complete"},
            {"message": "Detalhe as etapas clínicas e o acompanhamento com datas de calendário a partir do evento de referência. Identifique a fonte e suas limitações.",
             "expected_status": "complete"},
        ]
        metadata = {k: r[k] for k in METADATA_FIELDS}
        plan.append({"execution_id": r["execution_id"], "metadata": metadata, "context": context, "turns": turns,
                     "review_only": {"condition": conditions[injury]["name"], "triage": triage,
                                     "full_case": c["caso_completo"], "gold_id": gold[(c["case_id"], g)].get("gold_id")}})
    if not plan:
        raise ValueError("No planned conversations")
    plan.sort(key=lambda p: p["metadata"]["ordem_execucao"])
    plan_text = "".join(encode(p) + "\n" for p in plan)
    references_text = encode(gold_rows)
    manifest = {"schema": SCHEMA, "created_at_utc": now(), "reviewed_by": reviewed_by,
                "flow": FLOW, "clock": "Controlled evaluation_date from the workbook; not a form input",
                "version_meaning": "1=all mapped triage selections; 2=withheld mapped selections. Additional clinical facts arrive in chat turn 2.",
                "interface_sha256": digest((SOURCE_ROOT / "app/static/index.html").read_bytes()),
                "catalog_sha256": digest((SOURCE_ROOT / "knowledge/catalog.json").read_bytes()),
                "review_note": review_note, "source_name": source.name,
                "workbook_sha256": digest(source.read_bytes()), "plan_sha256": digest(plan_text.encode()),
                "references_sha256": digest(references_text.encode()),
                "n_cases": len(cases), "n_families": len({c['template_id'] for c in case_rows}),
                "n_conversations": len(plan), "n_turns": sum(len(p['turns']) for p in plan),
                "result_headers": result_headers,
                "review_scope": "Attestation by the named reviewer; not a claim of two independent reviews"}
    destination.mkdir(parents=True, exist_ok=True)
    (destination / "plan.jsonl").write_text(plan_text, encoding="utf-8")
    (destination / "references.json").write_text(references_text, encoding="utf-8")
    (destination / "manifest.json").write_text(encode(manifest), encoding="utf-8")
    return manifest


def load_plan(directory):
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    raw = (directory / "plan.jsonl").read_bytes()
    refs = (directory / "references.json").read_bytes()
    if manifest["schema"] != SCHEMA or digest(raw) != manifest["plan_sha256"] or digest(refs) != manifest["references_sha256"]:
        raise ValueError("The frozen plan/reference files changed; prepare a separate study version")
    return manifest, [json.loads(line) for line in raw.decode().splitlines() if line.strip()]


@contextmanager
def lock(directory):
    with (directory / ".collection.lock").open("a+b") as handle:
        if os.name == "nt":
            import msvcrt
            handle.seek(0); handle.write(b"0"); handle.flush(); handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
        try:
            yield
        finally:
            if os.name == "nt":
                handle.seek(0); msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle, fcntl.LOCK_UN)


def database(directory):
    db = sqlite3.connect(directory / "collection.sqlite3")
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("PRAGMA synchronous=FULL")
    db.execute("CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
    db.execute("""CREATE TABLE IF NOT EXISTS attempts (
        id INTEGER PRIMARY KEY, execution_id TEXT NOT NULL, turn INTEGER NOT NULL,
        state TEXT NOT NULL, started_at TEXT NOT NULL, request TEXT NOT NULL,
        reply TEXT, audit TEXT, error TEXT, latency_s REAL)""")
    db.commit()
    return db


def bind(db, key, value):
    serialized = encode(value)
    previous = db.execute("SELECT value FROM metadata WHERE key=?", (key,)).fetchone()
    if previous and previous[0] != serialized:
        raise ValueError(f"Collection {key} changed; use a new directory/version")
    db.execute("INSERT OR IGNORE INTO metadata VALUES (?,?)", (key, serialized)); db.commit()


def audit_matches(audit, server):
    keys = ("configured_model", "knowledge_sha256", "catalog_sha256", "service_sha256", "commit_git", "model_parameters")
    return all(audit.get(k) == server.get(k) for k in keys) and bool(audit.get("model")) and bool(audit.get("prompt_sha256"))


def collect(directory, url, max_conversations=1, retry_failed=False, retry_uncertain=False, client=None):
    if max_conversations < 1:
        raise ValueError("max_conversations must be positive")
    manifest, plan = load_plan(directory)
    if not url.startswith(("https://", "http://127.0.0.1:", "http://localhost:")):
        raise ValueError("Use HTTPS except for a local test server")
    owned = client is None
    client = client or httpx.Client(timeout=110)
    attempted = 0
    try:
        with lock(directory):
            db = database(directory)
            try:
                bind(db, "plan_sha256", manifest["plan_sha256"])
                bind(db, "references_sha256", manifest["references_sha256"])
                bind(db, "workbook_sha256", manifest["workbook_sha256"])
                bind(db, "collector", {"sha256": digest(Path(__file__).read_bytes()),
                     "python": platform.python_version(), "httpx": version("httpx"), "openpyxl": version("openpyxl")})
                info = client.get(url.rstrip("/") + "/api/evaluation-info")
                info.raise_for_status(); server = info.json()
                required = ("configured_model", "knowledge_sha256", "catalog_sha256", "service_sha256", "model_parameters")
                if not isinstance(server, dict) or server.get("evaluation_contract") != SCHEMA or not all(server.get(k) for k in required):
                    raise ValueError("Deploy the evaluation API update before collecting paid responses")
                bind(db, "server", server); bind(db, "url", url.rstrip("/"))
                for case in plan:
                    rows = db.execute("SELECT * FROM attempts WHERE execution_id=? ORDER BY id", (case["execution_id"],)).fetchall()
                    successful = {r["turn"]: r for r in rows if r["state"] == "success"}
                    if len(successful) == len(case["turns"]):
                        continue
                    latest = rows[-1] if rows else None
                    if latest and latest["state"] == "started" and not retry_uncertain:
                        raise ValueError("Interrupted request with unknown billing/outcome. Inspect attempts; use --retry-uncertain explicitly to resend.")
                    if latest and latest["state"] == "error" and not retry_failed:
                        continue
                    if attempted >= max_conversations:
                        break
                    attempted += 1
                    history, context = [], dict(case["context"])
                    for number, turn in enumerate(case["turns"], 1):
                        if number in successful:
                            reply = json.loads(successful[number]["reply"])
                        else:
                            payload = {"context": context, "history": history, "message": turn["message"]}
                            cursor = db.execute("INSERT INTO attempts (execution_id,turn,state,started_at,request) VALUES (?,?,?,?,?)",
                                                (case["execution_id"], number, "started", now(), encode(payload)))
                            db.commit(); attempt_id = cursor.lastrowid; start = time.monotonic()
                            reply = audit = None
                            fatal = False
                            try:
                                response = client.post(url.rstrip("/") + "/api/chat", json=payload)
                                response.raise_for_status(); reply = response.json()
                                if not isinstance(reply, dict) or not isinstance(reply.get("message"), str) or not reply["message"].strip() or reply.get("guideline") not in ("ESE", "AAE", "COMPARE"):
                                    raise ValueError("Malformed conversational response")
                                if reply.get("assessment") is not None and not isinstance(reply["assessment"], dict):
                                    raise ValueError("Malformed structured assessment")
                                audit = json.loads(response.headers.get("x-dart-audit", "null"))
                                if not isinstance(audit, dict) or not audit_matches(audit, server):
                                    fatal = True
                                    raise ValueError("Server/model configuration or audit contract changed during collection")
                                try:
                                    bind(db, "reported_model", audit["model"])
                                except ValueError:
                                    fatal = True
                                    raise
                                db.execute("UPDATE attempts SET state='success', reply=?,audit=?,latency_s=? WHERE id=?",
                                           (encode(reply), encode(audit), time.monotonic()-start, attempt_id)); db.commit()
                            except (httpx.HTTPError, KeyError, TypeError, ValueError) as exc:
                                db.execute("UPDATE attempts SET state='error',reply=?,audit=?,error=?,latency_s=? WHERE id=?",
                                           (encode(reply) if reply is not None else None, encode(audit) if audit is not None else None,
                                            str(exc), time.monotonic()-start, attempt_id)); db.commit()
                                if fatal:
                                    raise
                                break
                        history.extend([{"role": "user", "content": turn["message"]},
                                        {"role": "assistant", "content": encode(reply)}])
                        context["guideline"] = reply["guideline"]
                    print(f"Collected {case['execution_id']}", flush=True)
            finally:
                db.close()
    finally:
        if owned:
            client.close()
        export(directory)
    return attempted


def safe_csv(value):
    # Keep model-generated strings from being interpreted as spreadsheet formulas.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export(directory):
    with lock(directory):
        return _export(directory)


def _export(directory):
    manifest, plan = load_plan(directory)
    db = database(directory)
    try:
        attempts = [dict(r) for r in db.execute("SELECT * FROM attempts ORDER BY id")]
        for a in attempts:
            for key in ("request", "reply", "audit"):
                a[key] = json.loads(a[key]) if a[key] is not None else None
        grouped = {}
        for a in attempts:
            grouped.setdefault(a["execution_id"], []).append(a)
        metadata = {r["key"]: json.loads(r["value"]) for r in db.execute("SELECT * FROM metadata")}
        (directory / "collection-metadata.json").write_text(encode(metadata), encoding="utf-8")
        extra = ["configured_model", "service_sha256", "workbook_sha256", "collection_state", "attempts_total",
                 "status_t1", "status_t2", "status_t3", "status_match_t1", "status_match_t2", "status_match_t3"]
        output = directory / "resultados.csv"
        temporary = output.with_suffix(".csv.tmp")
        raw = directory / "transcricoes.jsonl"
        with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=manifest["result_headers"] + extra, delimiter=";")
            writer.writeheader()
            for case in sorted(plan, key=lambda c: c["metadata"]["ordem_execucao"]):
                observed = grouped.get(case["execution_id"], [])
                success = {a["turn"]: a for a in observed if a["state"] == "success"}
                complete = len(success) == len(case["turns"])
                state = "complete" if complete else "uncertain" if observed and observed[-1]["state"] == "started" else "error" if observed else "planned"
                row = {**case["metadata"], "status_execucao": 1 if complete else 2 if observed else 0,
                       "collection_state": state, "attempts_total": len(observed), "workbook_sha256": manifest["workbook_sha256"]}
                if observed:
                    row["executado_em_utc"] = observed[0]["started_at"]
                    if all(a["latency_s"] is not None for a in observed):
                        row["latencia_s"] = sum(a["latency_s"] for a in observed)
                    row["erro_tecnico"] = " | ".join(a["error"] or ("Interrupted request; outcome unknown" if a["state"] == "started" else "") for a in observed if a["error"] or a["state"] == "started")
                    row["transcricao_arquivo"] = raw.name + "#" + case["execution_id"]
                    audits = [a["audit"] for a in observed if a["audit"]]
                    if audits:
                        audit = audits[-1]
                        row.update(modelo=audit.get("model"), configured_model=audit.get("configured_model"),
                                   prompt_sha256=audit.get("prompt_sha256"), catalogo_sha256=audit.get("catalog_sha256"),
                                   commit_git=audit.get("commit_git"), service_sha256=audit.get("service_sha256"))
                        if len(audits) == len(observed) and all(isinstance(a.get("usage", {}).get("total_tokens"), int) for a in audits):
                            row["tokens_total"] = sum(a["usage"]["total_tokens"] for a in audits)
                    # Costs and clinical scores remain blank until actually assessed.
                for number, a in success.items():
                    row[f"resposta_t{number}"] = a["reply"]["message"]
                    assessment = a["reply"].get("assessment") or {}
                    status = assessment.get("status", "no_assessment")
                    row[f"status_t{number}"] = status
                    expected = case["turns"][number-1].get("expected_status")
                    if expected is not None:
                        row[f"status_match_t{number}"] = int(status == expected)
                writer.writerow({k: safe_csv(v) for k, v in row.items()})
        os.replace(temporary, output)
        raw_tmp = raw.with_suffix(".jsonl.tmp")
        raw_tmp.write_text("".join(encode(a)+"\n" for a in attempts), encoding="utf-8")
        os.replace(raw_tmp, raw)
        export_comparison(directory, manifest, plan, grouped)
        return output
    finally:
        db.close()


def export_comparison(directory, manifest, plan, grouped):
    """One observed reply per row, with actual inputs and a separate, unsent rubric."""
    refs = json.loads((directory / "references.json").read_text(encoding="utf-8"))
    gold = {(g.get("case_id"), g.get("diretriz_code")): g for g in refs}
    fields = ["execution_id", "case_id", "versao_code", "diretriz", "repeticao", "turno", "fluxo",
              "opcoes_da_triagem", "campos_omitidos_intencionalmente", "dados_clinicos_disponiveis",
              "mensagem_enviada", "resposta_exata_api", "resposta_exibida_no_site", "criterios_do_gabarito",
              "correto", "completo", "erro_critico", "justificativa", "data_evento", "relogio_avaliacao",
              "fonte_gabarito", "estado_tecnico", "historico_preservado"]
    out = directory / "conferencia.csv"
    temp = out.with_suffix(".csv.tmp")
    with temp.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, delimiter=";")
        writer.writeheader()
        for case in plan:
            observed = grouped.get(case["execution_id"], [])
            success = {a["turn"]: a for a in observed if a["state"] == "success"}
            review = case.get("review_only", {})
            triage = review.get("triage", [])
            form = "\n".join(f"{q['key']} - {q['field']}: {q['selected_label']}" for q in triage)
            withheld = "\n".join(f"{q['key']} - {q['field']} (valor retido: {q['reference_label']})" for q in triage if q["withheld"])
            ref = gold.get((case["metadata"]["case_id"], case["metadata"]["diretriz_code"]), {})
            rubric = "\n\n".join(f"{k}: {ref[k]}" for k in ("decisao_esperada", "abordagem_pulpar", "etapas_clinicas", "acompanhamento_datas", "transparencia_fonte", "seguranca") if ref.get(k))
            for turn, a in sorted(success.items()):
                prompt = a["request"]
                first = ("Avaliar se solicita os dados pertinentes ainda ausentes, sem presumir achados. "
                         "Formulário preenchido não garante que todo dado clínico necessário esteja disponível. "
                         "Uma orientação condicional segura é aceitável; não exigir conduta definitiva antes dos dados necessários.")
                available = prompt["context"].get("case_text", "")
                # Include only clinical messages already delivered, never future/reference facts.
                clinical_messages = [m["content"] for m in prompt.get("history", []) if m["role"] == "user" and m["content"] != initial_message()]
                if turn > 1:
                    clinical_messages.append(prompt["message"])
                available = "\n\n".join(x for x in [available, *clinical_messages] if x)
                row = {"execution_id": case["execution_id"], "case_id": case["metadata"]["case_id"],
                       "versao_code": case["metadata"]["versao_code"], "diretriz": prompt["context"]["guideline"],
                       "repeticao": case["metadata"]["repeticao"], "turno": turn,
                       "fluxo": manifest.get("flow", "legacy-api-narrative-v1"), "opcoes_da_triagem": form,
                       "campos_omitidos_intencionalmente": withheld, "dados_clinicos_disponiveis": available,
                       "mensagem_enviada": prompt["message"], "resposta_exata_api": a["reply"]["message"],
                       "resposta_exibida_no_site": a["reply"]["message"].replace("**", ""),
                       "criterios_do_gabarito": first if turn == 1 else rubric,
                       "data_evento": prompt["context"].get("reference_date"),
                       "relogio_avaliacao": prompt["context"].get("evaluation_date"),
                       "fonte_gabarito": f"{ref.get('gold_id', '')}; DOI {ref.get('doi', '')}; PDF pp. {ref.get('paginas_pdf', '')}",
                       "estado_tecnico": a["state"], "historico_preservado": int(bool(prompt.get("history")))}
                writer.writerow({k: safe_csv(v) for k, v in row.items()})
    os.replace(temp, out)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("prepare", help="Read Excel and freeze a collection plan; no model calls")
    p.add_argument("workbook", type=Path); p.add_argument("--out-dir", type=Path, required=True)
    p.add_argument("--reviewed-by", required=True); p.add_argument("--review-note", default="")
    p = sub.add_parser("run", help="Collect bounded conversations; incurs server API usage")
    p.add_argument("directory", type=Path); p.add_argument("--url", required=True)
    p.add_argument("--max-conversations", type=int, default=1)
    p.add_argument("--retry-failed", action="store_true"); p.add_argument("--retry-uncertain", action="store_true")
    p = sub.add_parser("export", help="Regenerate Excel-readable CSV and raw transcripts; no model calls")
    p.add_argument("directory", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "prepare":
            m = prepare(args.workbook, args.out_dir, args.reviewed_by, args.review_note)
            export(args.out_dir)
            print(f"Prepared {m['n_cases']} cases / {m['n_conversations']} conversations. No API calls.")
        elif args.command == "run":
            collect(args.directory, args.url, args.max_conversations, args.retry_failed, args.retry_uncertain)
        else:
            print(export(args.directory))
    except (ValueError, OSError, httpx.HTTPError, sqlite3.Error) as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
